from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_AUTO_SIZE

OUT = '/Users/kurniawanrizki/wa-autoreply-umkm/BalasWA_Business_Model_Canvas.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

INK = RGBColor(25, 21, 32); MUTED = RGBColor(104, 96, 111); PURPLE = RGBColor(109, 69, 215)
LAV = RGBColor(244, 240, 251); LAV2 = RGBColor(235, 226, 250); WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(222, 214, 232); GREEN = RGBColor(22, 128, 92); YELLOW = RGBColor(255, 243, 218)

def bg(slide, color=WHITE):
    shape=slide.background.fill; shape.solid(); shape.fore_color.rgb=color

def box(slide,x,y,w,h,fill=WHITE,line=LINE,radius=True):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    sh.line.color.rgb=line; sh.line.width=Pt(.65)
    return sh

def text(slide,x,y,w,h,txt,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,margin=.05):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin)
    p=tf.paragraphs[0]; p.alignment=align
    for i,line in enumerate(txt.split('\n')):
        if i: p=tf.add_paragraph(); p.alignment=align
        r=p.add_run(); r.text=line; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
        p.space_after=Pt(0)
    return tb

def title(slide,kicker,heading,sub=''):
    text(slide,.62,.38,12,.24,kicker.upper(),8,PURPLE,True)
    text(slide,.62,.68,12,.55,heading,27,INK,True)
    if sub:text(slide,.62,1.28,12,.28,sub,10,MUTED)

def footer(slide,n):
    text(slide,.62,7.13,5,.16,'BALASWA  •  DRAFT STRATEGI UMKM',7,MUTED,True)
    text(slide,12.1,7.13,.6,.16,str(n).zfill(2),7,MUTED,True,PP_ALIGN.RIGHT)

# 1
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,LAV)
box(s,.62,.62,1.0,.35,PURPLE,PURPLE); text(s,.72,.70,.8,.15,'BALASWA',8,WHITE,True,PP_ALIGN.CENTER)
text(s,.62,1.55,9.8,1.25,'Business Model\nCanvas',40,INK,True)
text(s,.66,3.06,7.7,.5,'SaaS balasan WhatsApp berbantuan AI untuk UMKM Indonesia.',16,MUTED)
text(s,.66,3.72,6.5,.4,'Model: trial 7 hari → langganan bulanan.',11,PURPLE,True)
# visual blocks
for x,y,w,h,c in [(9.45,1.05,2.72,1.2,PURPLE),(10.32,2.48,2.05,.9,LAV2),(8.82,3.6,2.55,1.1,WHITE),(10.58,5.02,1.8,.7,LAV2)]: box(s,x,y,w,h,c,c)
text(s,9.72,1.40,2.1,.3,'CHAT MASUK',12,WHITE,True,PP_ALIGN.CENTER)
text(s,10.53,2.77,1.6,.25,'AI MENJAWAB',10,PURPLE,True,PP_ALIGN.CENTER)
text(s,9.08,4.0,2.0,.25,'UMKM TETAP MELAYANI',11,INK,True,PP_ALIGN.CENTER)
footer(s,1)

# BMC 2
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,WHITE); title(s,'Model bisnis','Business Model Canvas — BalasWA','Asumsi awal. Validasi dengan 10–15 UMKM pada fase pilot.'); footer(s,2)
# canvas geometry
x0=.62; y0=1.78; gap=.07; col=2.36; htop=1.95; hbot=1.95
items=[
 ('Mitra Kunci','• Penyedia AI / 9Router\n• Komunitas pengembang\n• Hosting & pembayaran\n• Komunitas UMKM\n• Reseller / agensi digital',x0,y0,col,htop+hbot+gap),
 ('Aktivitas Kunci','• Menjaga sesi WhatsApp\n• Balas chat relevan\n• Sinkron profil & katalog\n• Dukungan onboarding\n• Kontrol biaya AI',x0+col+gap,y0,col,htop),
 ('Sumber Daya Kunci','• Platform multi-tenant\n• Mesin aturan balasan\n• Integrasi AI\n• Data profil & stok\n• Dukungan pelanggan',x0+col+gap,y0+htop+gap,col,hbot),
 ('Nilai yang Ditawarkan','• Balas pelanggan lebih cepat\n• Jawaban berdasar katalog & stok\n• Pemilik tetap pegang nomor WA\n• Tanpa teknis: scan QR, isi profil\n• Hemat waktu admin',x0+2*(col+gap),y0,col,htop+hbot+gap),
 ('Hubungan Pelanggan','• Coba sendiri via trial\n• Onboarding checklist\n• Pusat bantuan singkat\n• Dukungan chat untuk masalah akun\n• Upgrade mandiri',x0+3*(col+gap),y0,col,htop),
 ('Saluran','• Landing page BalasWA\n• Konten edukasi UMKM\n• Komunitas / webinar\n• Mitra agensi\n• Rekomendasi pengguna',x0+3*(col+gap),y0+htop+gap,col,hbot),
 ('Segmen Pelanggan','• UMKM yang jualan via WhatsApp\n• Toko fashion, kuliner, jasa\n• Admin terbatas / owner-operated\n• Pelanggan butuh respons di luar jam sibuk\n• Tahap awal: usaha dengan katalog sederhana',x0+4*(col+gap),y0,col,htop+hbot+gap),
]
for head,body,x,y,w,h in items:
    box(s,x,y,w,h,LAV if head in ['Nilai yang Ditawarkan','Segmen Pelanggan'] else WHITE, LINE)
    text(s,x+.13,y+.14,w-.26,.27,head.upper(),8,PURPLE,True)
    text(s,x+.13,y+.5,w-.26,h-.62,body,9,INK)
# bottom cost/revenue
box(s,x0,y0+htop+hbot+gap*2,6.13,1.0,WHITE,LINE); box(s,6.82,y0+htop+hbot+gap*2,5.89,1.0,WHITE,LINE)
text(s,.76,5.98,5.7,.2,'STRUKTUR BIAYA',8,PURPLE,True); text(s,.76,6.25,5.7,.35,'AI per balasan • server/Chrome • pengembangan • dukungan • akuisisi',10,INK)
text(s,6.96,5.98,5.3,.2,'ARUS PENDAPATAN',8,PURPLE,True); text(s,6.96,6.25,5.3,.35,'Starter Rp99rb/bulan • Tumbuh Rp199rb/bulan • add-on kuota (usulan)',10,INK)

# 3
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,WHITE); title(s,'Ekonomi unit','Harga, batas pakai, dan hipotesis yang perlu diuji','Harga final mengikuti biaya AI aktual dan tingkat penggunaan pelanggan.'); footer(s,3)
# pricing
for x,name,price,limit,accent in [(0.72,'TRIAL','Rp0','7 hari • 100 balasan',LAV2),(4.42,'STARTER','Rp99.000','3.000 balasan / bulan',PURPLE),(8.12,'TUMBUH','Rp199.000','10.000 balasan / bulan',LAV2)]:
 box(s,x,1.95,3.15,2.05,accent,accent)
 c=WHITE if accent==PURPLE else INK; mc=WHITE if accent==PURPLE else MUTED
 text(s,x+.22,2.2,2.7,.2,name,9,c,True); text(s,x+.22,2.62,2.7,.42,price,25,c,True); text(s,x+.22,3.35,2.7,.2,limit,10,mc,True)
# metrics / validation
text(s,.72,4.55,5.4,.25,'METRIK UTAMA',10,PURPLE,True)
for i,(h,b) in enumerate([('Aktivasi','Profil + 1 produk + WA tersambung'),('Nilai awal','Balasan pertama terkirim'),('Retensi','Usaha masih aktif minggu ke-4'),('Beban AI','Biaya AI ÷ pendapatan pelanggan')]):
 y=4.92+i*.45; text(s,.75,y,1.55,.18,h,10,INK,True); text(s,2.35,y,3.8,.18,b,10,MUTED)
box(s,7.0,4.48,5.64,1.92,YELLOW,YELLOW); text(s,7.28,4.78,5.1,.2,'PILOT 30 HARI',10,INK,True)
text(s,7.28,5.15,5.0,.9,'Rekrut 10–15 UMKM. Ukur: waktu respons sebelum/sesudah, persentase chat relevan, biaya AI per akun, serta kemauan membayar setelah trial.',13,INK)

text(s,.72,6.65,11.9,.3,'Risiko: whatsapp-web.js bukan API resmi; sesi dapat terputus atau akun dibatasi. Mitra dan harga masih rencana.',9,MUTED)
prs.save(OUT)
print(OUT)
