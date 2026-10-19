from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).parent
OUT = ROOT / 'BalasWA_Business_Model_Canvas_12_Slide.pptx'
p = Presentation()
p.slide_width, p.slide_height = Inches(16), Inches(9)
INK, MUTED, PURPLE, LIGHT, LINE = '211A30', '70667D', '6D45D7', 'F5F1FB', 'DDD4E9'

def rect(s,x,y,w,h,color):
    a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb=RGBColor.from_string(color); a.line.fill.background()

def txt(s,x,y,w,h,t,size=22,color=INK,bold=False):
    a=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f=a.text_frame; f.word_wrap=True
    f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
    for i,line in enumerate(t.split('\n')):
        q=f.paragraphs[0] if not i else f.add_paragraph()
        q.text=line; q.font.name='Aptos'; q.font.size=Pt(size)
        q.font.bold=bold; q.font.color.rgb=RGBColor.from_string(color)
        q.space_after=Pt(10)
    return a

def slide(k,title,sub):
    s=p.slides.add_slide(p.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string('FFFFFF')
    txt(s,.7,.4,14,.3,k.upper(),12,PURPLE,True)
    txt(s,.7,.95,14.6,.75,title,34,INK,True)
    txt(s,.7,1.85,14.5,.65,sub,18,MUTED)
    rect(s,.7,8.38,14.6,.015,LINE)
    txt(s,.7,8.58,13,.22,'BALASWA   /   BUSINESS MODEL CANVAS   /   RANCANGAN BISNIS',10,MUTED)
    txt(s,14.5,8.53,.8,.3,f'{len(p.slides):02d} / 12',10,MUTED)
    return s

def columns(s,left_title,left,right_title,right,bottom=''):
    rect(s,8.2,2.85,7.1,4.6,LIGHT)
    txt(s,.8,2.95,6.8,.45,left_title,23,PURPLE,True)
    txt(s,.8,3.65,6.7,3.65,'\n'.join(left),22)
    txt(s,8.55,3.05,6.35,.45,right_title,23,PURPLE,True)
    txt(s,8.55,3.75,6.25,3.5,'\n'.join(right),22)
    if bottom: txt(s,.8,7.75,14.2,.4,bottom,16,MUTED)

s=slide('01 / Konsep bisnis','BalasWA','Layanan berlangganan untuk membantu UMKM membalas pelanggan melalui WhatsApp.')
txt(s,.8,3,11,2.1,'Pelanggan bertanya.\nUsaha tetap melayani.',48,INK,True)
txt(s,.8,5.65,11.7,1.2,'Balasan berbantuan AI berdasarkan profil usaha, katalog, harga, stok, dan instruksi pemilik.',26,MUTED)
rect(s,.8,7.25,10.8,.65,LIGHT)
txt(s,1,7.42,10.4,.3,'Fokus awal: UMKM dengan satu nomor WhatsApp dan katalog sederhana.',17,PURPLE,True)

s=slide('02 / Masalah & peluang','Chat berulang menyita waktu pemilik','Hipotesis masalah pelanggan; belum merupakan hasil survei atau klaim dampak terukur.')
columns(s,'Situasi yang ingin diselesaikan',['Pertanyaan harga, stok, lokasi, dan jam buka berulang.','Pemilik membagi waktu antara operasional dan chat.','Jawaban terlambat saat toko sibuk atau admin tidak tersedia.'],'Peluang BalasWA',['Menjawab pertanyaan rutin dari data usaha.','Mengurangi pengetikan jawaban yang sama.','Membantu layanan awal tanpa mengganti seluruh peran admin.'],'Ukuran keberhasilan: waktu respons, ketepatan jawaban, dan waktu admin yang dihemat.')

s=slide('03 / Nilai yang ditawarkan','Bukan sekadar chatbot dengan prompt','Pemilik menyiapkan fakta usaha; AI membantu merangkai jawaban yang relevan.')
columns(s,'Manfaat untuk UMKM',['Pengaturan lewat web, tanpa terminal.','Nomor usaha ditautkan melalui scan QR.','Jawaban memakai profil, produk, harga, dan stok yang diisi pemilik.'],'Batas layanan yang jelas',['Stok harus diperbarui pemilik; bukan bukti stok fisik secara otomatis.','Bot tidak mengonfirmasi pembayaran atau pesanan tanpa integrasi.','Kasus rumit tetap membutuhkan admin manusia.'],'Instruksi tambahan mengatur gaya bahasa; fakta bisnis tetap berasal dari data terstruktur.')

s=slide('04 / Peta model bisnis','Business Model Canvas — sembilan komponen','Ringkasan strategi; mitra, saluran, dan pendapatan merupakan rancangan yang perlu diuji.')
blocks=[('Mitra kunci','Penyedia AI dan hosting\nKomunitas UMKM\nAgensi digital',.7,2.8,2.85,3.7),('Aktivitas kunci','Menjaga koneksi WA\nMemperbaiki jawaban\nMendampingi pengguna',3.65,2.8,2.85,1.8),('Sumber daya kunci','Platform dan tim\nData usaha\nInfrastruktur sesi WA',3.65,4.7,2.85,1.8),('Nilai yang ditawarkan','Respons lebih cepat\nJawaban berbasis katalog\nPengaturan mandiri\nMenghemat waktu admin',6.6,2.8,2.8,3.7),('Hubungan pelanggan','Trial mandiri\nPanduan aktivasi\nBantuan penggunaan',9.5,2.8,2.85,1.8),('Saluran','Web BalasWA\nKonten edukasi\nKomunitas dan rujukan',9.5,4.7,2.85,1.8),('Segmen pelanggan','UMKM via WhatsApp\nFashion, kuliner, jasa\nPemilik dengan admin terbatas',12.45,2.8,2.85,3.7),('Struktur biaya','AI • server dan Chrome • pengembangan • dukungan • pemasaran',.7,6.6,7.25,1.35),('Arus pendapatan','Langganan Starter / Tumbuh; tambahan kuota masih usulan',8.05,6.6,7.25,1.35)]
for h,b,x,y,w,hh in blocks:
    rect(s,x,y,w,hh,LIGHT)
    txt(s,x+.15,y+.15,w-.3,.4,h,16,PURPLE,True)
    txt(s,x+.15,y+.65,w-.3,hh-.7,b,15)

s=slide('05 / Segmen pelanggan','Mulai dari usaha yang kebutuhannya sederhana','Pembeli layanan: pemilik usaha. Pengguna harian: pemilik atau admin toko.')
columns(s,'Sasaran awal',['Toko fashion dengan katalog dan harga jelas.','Kuliner dengan menu, jam buka, serta area pengiriman tetap.','Usaha jasa dengan pertanyaan awal yang berulang.'],'Kriteria pilot',['Menggunakan satu nomor usaha aktif.','Bersedia mengisi dan memperbarui data produk.','Mau menilai jawaban bot dan membagikan kendala penggunaan.'],'Belum menjadi prioritas: perusahaan dengan banyak cabang, banyak admin, atau alur transaksi kompleks.')

s=slide('06 / Saluran & hubungan pelanggan','Dari mengenal produk sampai berlangganan','Rencana pemasaran awal: demonstrasi nyata, bukan janji otomatisasi seluruh bisnis.')
columns(s,'Saluran akuisisi',['Web dengan contoh pertanyaan dan jawaban.','Konten singkat tentang layanan pelanggan UMKM.','Demo melalui komunitas usaha dan mitra agensi.'],'Hubungan pelanggan',['Trial 7 hari untuk mencoba manfaat.','Panduan: isi profil, tambah produk, tautkan WhatsApp.','Bantuan aktivasi dan evaluasi sebelum trial berakhir.'],'Pantau asal pendaftaran, keberhasilan aktivasi, konversi berbayar, dan alasan berhenti.')

s=slide('07 / Aktivitas kunci','Keandalan layanan menjadi pekerjaan utama','Produk bernilai hanya jika pesan diterima, jawaban tepat, dan balasan benar-benar terkirim.')
columns(s,'Operasional produk',['Menjaga sesi WhatsApp dan pemulihan koneksi.','Menyaring pesan yang tidak relevan.','Menguji ketepatan jawaban dan membatasi penggunaan AI.'],'Operasional bisnis',['Mendampingi aktivasi pengguna baru.','Menangani keluhan serta kegagalan balasan.','Menguji harga dan meninjau biaya per pelanggan.'],'Log harus membedakan pesan masuk, pemrosesan AI, keberhasilan kirim, dan kegagalan.')

s=slide('08 / Sumber daya & mitra','Teknologi mendukung; data usaha menentukan','Daftar mitra berikut adalah calon penyedia atau kanal kerja sama, bukan kemitraan yang sudah disepakati.')
columns(s,'Sumber daya utama',['Aplikasi web, akun pengguna, dan penyimpanan sesi.','Integrasi whatsapp-web.js, 9Router, serta penyedia AI.','Tim pengembangan dan dukungan; data usaha yang akurat.'],'Calon mitra',['Penyedia hosting dan model AI.','Komunitas UMKM serta agensi digital.','Penyedia pembayaran untuk tahap penagihan berikutnya.'],'WhatsApp adalah platform yang menjadi ketergantungan produk; bukan mitra resmi BalasWA.')

s=slide('09 / Arus pendapatan','Langganan bulanan dengan batas penggunaan','Rancangan harga dari PRD; kelayakan biaya dan kesediaan membayar belum tervalidasi.')
for x,name,price,limit in [(.8,'TRIAL','Gratis','7 hari • 100 balasan AI'),(5.7,'STARTER','Rp99.000','3.000 balasan AI / bulan'),(10.6,'TUMBUH','Rp199.000','10.000 balasan AI / bulan')]:
    rect(s,x,3,4.6,3.35,LIGHT)
    txt(s,x+.25,3.3,4.1,.4,name,18,PURPLE,True)
    txt(s,x+.25,4.05,4.1,.7,price,34,INK,True)
    txt(s,x+.25,5,4.1,.4,'per bulan' if name!='TRIAL' else 'tanpa biaya trial',18,MUTED)
    txt(s,x+.25,5.65,4.1,.45,limit,17)
txt(s,.8,6.95,14,1,'Tambahan kuota: usulan tahap berikutnya. Integrasi pembayaran otomatis belum termasuk klaim fitur yang sudah tersedia.',21,MUTED)

s=slide('10 / Struktur biaya','Jangan menilai margin dari biaya AI saja','Belum ada proyeksi laba karena data biaya aktual dan pola penggunaan pilot belum tersedia.')
columns(s,'Komponen biaya',['Variabel: token AI, sesi Chrome, dan dukungan per akun.','Tetap: pengembangan, pemantauan, serta infrastruktur dasar.','Akuisisi: konten, demo, mitra, dan promosi.'],'Cara mengukur kelayakan',['Catat biaya AI dan server per akun aktif.','Margin kontribusi = pendapatan dikurangi biaya langsung layanan.','Titik impas akun = biaya tetap dibagi margin kontribusi per akun.'],'Kuota dan harga ditinjau setelah pola penggunaan nyata terkumpul; bukan berdasarkan asumsi token saja.')

s=slide('11 / Risiko & mitigasi','Kepercayaan pelanggan harus dijaga','whatsapp-web.js bukan API resmi WhatsApp; koneksi dapat berubah dan akun dapat terkena pembatasan.')
columns(s,'Risiko utama',['Sesi terputus atau perubahan WhatsApp Web.','Jawaban keliru akibat data usang atau AI.','Kebocoran percakapan, akses akun, dan kenaikan biaya.'],'Mitigasi yang direncanakan',['Status koneksi jelas; panduan penautan ulang; evaluasi API resmi.','Data terstruktur, evaluasi jawaban, dan batas kewenangan bot.','Pemisahan akun, pembatasan akses, kebijakan retensi, dan kuota.'],'Mitigasi perlu diuji. Hindari janji bebas blokir, selalu tersedia, atau jawaban selalu benar.')

s=slide('12 / Rencana validasi','Uji manfaat sebelum memperbesar bisnis','Usulan pilot 30 hari dengan 10–15 UMKM; ini target eksperimen, bukan jumlah pelanggan aktual.')
columns(s,'Tahap pelaksanaan',['Minggu 1: rekrut peserta dan ukur kondisi awal.','Minggu 2–3: aktivasi, amati pemakaian, perbaiki kendala.','Minggu 4: wawancara, tawarkan paket, evaluasi kelanjutan.'],'Bukti yang dikumpulkan',['Waktu respons serta ketepatan jawaban.','Aktivasi berhasil, pemakaian berulang, dan waktu admin.','Biaya per akun, kesediaan membayar, serta alasan berhenti.'],'Keputusan: lanjutkan jika manfaat, keandalan, dan margin terbukti; revisi harga atau sasaran jika belum.')

assert len(p.slides)==12
for s in p.slides:
    for a in s.shapes:
        assert a.left >= 0 and a.top >= 0
        assert a.left+a.width <= p.slide_width and a.top+a.height <= p.slide_height
p.save(OUT)
r=Presentation(OUT)
assert len(r.slides)==12
assert all(any(a.has_text_frame and a.text.strip() for a in s.shapes) for s in r.slides)
print(f'PASS: {len(r.slides)} slides; shape bounds and PPTX read-back verified. {OUT}')
